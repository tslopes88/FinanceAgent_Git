# -*- coding: utf-8 -*-
"""
FinanceAgent Desktop - Executive Dark Edition
Plataforma de Conciliação OFX & Inteligência Analítica Financeira
Arquitetura: 100% Desktop Nativo (CustomTkinter + Matplotlib TkAgg + SQLite3)
"""

import os
import re
import sys
import shutil
import sqlite3
import subprocess
from datetime import date, datetime
from pathlib import Path
import pandas as pd
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import customtkinter as ctk

# Matplotlib para Renderização de Gráficos Nativos no Tkinter
import matplotlib
matplotlib.use("TkAgg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg

# -------------------------------------------------------------
# 1. RESILIÊNCIA DE CAMINHOS & INDEPENDÊNCIA (.EXE / ONEDRIVE)
# -------------------------------------------------------------
def obter_diretorio_base() -> Path:
    """
    Retorna o diretório base real da aplicação.
    Suporta perfeitamente a execução via PyInstaller (sys.frozen) no OneDrive.
    """
    if getattr(sys, 'frozen', False):
        return Path(sys.executable).parent.resolve()
    else:
        return Path(__file__).parent.resolve()

BASE_DIR = obter_diretorio_base()
PASTA_ENTRADA = BASE_DIR / "entrada"
PASTA_SAIDA = BASE_DIR / "saida"
DB_PATH = BASE_DIR / "finanagent.db"
ARQUIVO_MIGRACAO_CSV = BASE_DIR / "extrato_consolidado_2026_08.csv"

PASTA_ENTRADA.mkdir(parents=True, exist_ok=True)
PASTA_SAIDA.mkdir(parents=True, exist_ok=True)

# -------------------------------------------------------------
# 2. MOTOR DE BANCO DE DADOS LOCAL (SQLITE3)
# -------------------------------------------------------------
def get_db_connection():
    conn = sqlite3.connect(str(DB_PATH), timeout=15.0)
    conn.row_factory = sqlite3.Row
    return conn

def sanitizar_texto(texto: str) -> str:
    t = str(texto).strip()
    while t.startswith(('=', '+', '-', '@', '\t', '\r')):
        t = t[1:].strip()
    return t

DADOS_DEMO_INICIAIS = [
    ("01/09/2026", "Extrato Bancário", "PIX RECEBIDO - CONSULTORIA TI CLIENTE ALFA", "Receitas Externas", "Consultoria", "Receita", 4850.00, "Compartilhado", "DEMO001"),
    ("02/09/2026", "Extrato Bancário", "TED RECEBIDA - PROJETO SOFTWARE BETA", "Receitas Externas", "Desenvolvimento", "Receita", 3200.00, "Compartilhado", "DEMO002"),
    ("03/09/2026", "Extrato Bancário", "RENDIMENTO APLICACAO AUTOMATICA CDI", "Receitas Externas", "Rendimentos", "Receita", 142.50, "Compartilhado", "DEMO003"),
    ("03/09/2026", "Cartão Débito", "POSTO SHELL COMBUSTIVEL", "Transporte & Mobilidade", "Combustível", "Despesa", -245.80, "Usuário Titular", "DEMO004"),
    ("04/09/2026", "Cartão Débito", "UBER *TRIP VIAGEM URBANA", "Transporte & Mobilidade", "Mobilidade", "Despesa", -38.50, "Usuário Titular", "DEMO005"),
    ("04/09/2026", "Cartão Débito", "SUPERMERCADO CARREFOUR COMPRAS MES", "Alimentação - Supermercados", "Mercado", "Despesa", -482.35, "Usuário Secundário", "DEMO006"),
    ("05/09/2026", "Cartão Débito", "HORTIFRUTI NATURAL DA TERRA", "Alimentação - Supermercados", "Hortifrúti", "Despesa", -198.40, "Usuário Secundário", "DEMO007"),
    ("05/09/2026", "Cartão Débito", "PADARIA CELEIRO DO PAO", "Alimentação - Padarias", "Café & Lanches", "Despesa", -45.90, "Usuário Titular", "DEMO008"),
    ("06/09/2026", "Cartão Débito", "IFOOD *RESTAURANTE BURGER", "Alimentação - Restaurantes", "Delivery", "Despesa", -112.00, "Compartilhado", "DEMO009"),
    ("06/09/2026", "Cartão Débito", "RESTAURANTE ESPETINHO GOURMET", "Alimentação - Restaurantes", "Almoço", "Despesa", -85.00, "Usuário Titular", "DEMO010"),
    ("07/09/2026", "Cartão Débito", "DROGASIL FARMACIA MEDICAMENTOS", "Saúde & Farmácia", "Medicamentos", "Despesa", -125.60, "Compartilhado", "DEMO011"),
    ("07/09/2026", "Débito Automático", "NETFLIX.COM MENSALIDADE", "Assinaturas & Apps", "Streaming", "Despesa", -55.90, "Usuário Titular", "DEMO012"),
    ("08/09/2026", "Débito Automático", "SPOTIFY ASSINATURA FAMILIAR", "Assinaturas & Apps", "Música", "Despesa", -34.90, "Compartilhado", "DEMO013"),
    ("08/09/2026", "Cartão Crédito", "MERCADO LIVRE ELETRONICOS", "Compras & Varejo", "Equipamentos", "Despesa", -289.90, "Usuário Titular", "DEMO014"),
    ("08/09/2026", "Cartão Crédito", "AMAZON BR MARKETPLACE", "Compras & Varejo", "Livros & Utilidades", "Despesa", -159.00, "Usuário Secundário", "DEMO015"),
    ("09/09/2026", "Tarifa Conta", "TARIFA MENSAL PACOTE SERVICOS", "Tarifas Bancárias", "Manutenção", "Despesa", -42.00, "Compartilhado", "DEMO016"),
    ("09/09/2026", "Imposto", "IOF OPERACOES FINANCEIRAS", "Tarifas Bancárias", "IOF", "Despesa", -8.45, "Compartilhado", "DEMO017"),
    ("09/09/2026", "PIX Transferência", "TRANSFERENCIA PIX REPASSE FAMILIAR", "Família & Pessoal", "Repasse Familiar", "Despesa", -650.00, "Usuário Secundário", "DEMO018"),
    ("10/09/2026", "Investimento", "APLICACAO COFRINHO RESERVA DE EMERGENCIA", "Investimentos & Reserva", "Reserva Financeira", "Reserva", -1000.00, "Compartilhado", "DEMO019"),
    ("10/09/2026", "Cartão Débito", "POSTO IPIRANGA ABASTECIMENTO", "Transporte & Mobilidade", "Combustível", "Despesa", -180.00, "Usuário Titular", "DEMO020")
]

def semear_dados_demonstracao_sqlite():
    with get_db_connection() as conn:
        registros_formatados = []
        for r in DADOS_DEMO_INICIAIS:
            registros_formatados.append((
                r[0], r[1], sanitizar_texto(r[2]), r[3], r[4], r[5], r[6], r[7], "Efetivado", r[8]
            ))
        conn.executemany("""
            INSERT INTO transacoes (data, origem, descricao, categoria, subcategoria, tipo_movimentacao, valor_brl, responsavel, status, fitid)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, registros_formatados)
        conn.commit()

def inicializar_banco():
    with get_db_connection() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS transacoes (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                data TEXT NOT NULL,
                origem TEXT,
                descricao TEXT NOT NULL,
                categoria TEXT NOT NULL,
                subcategoria TEXT,
                tipo_movimentacao TEXT NOT NULL,
                valor_brl REAL NOT NULL,
                responsavel TEXT NOT NULL,
                status TEXT DEFAULT 'Efetivado',
                fitid TEXT,
                criado_em TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)
        conn.commit()

        cursor = conn.cursor()
        cursor.execute("SELECT COUNT(*) FROM transacoes")
        total_registros = cursor.fetchone()[0]

        if total_registros == 0:
            if ARQUIVO_MIGRACAO_CSV.exists():
                try:
                    df_mig = pd.read_csv(ARQUIVO_MIGRACAO_CSV)
                    registros_mig = []
                    for _, r in df_mig.iterrows():
                        registros_mig.append((
                            str(r.get("Data", "01/08/2026")),
                            str(r.get("Origem", "Extrato Inicial")),
                            sanitizar_texto(r.get("Descricao", "Lancamento")),
                            str(r.get("Categoria", "Outros")),
                            str(r.get("Subcategoria", "")),
                            str(r.get("Tipo_Movimentacao", "Despesa")),
                            float(r.get("Valor_BRL", 0.0)),
                            str(r.get("Responsavel", "Compartilhado")),
                            "Efetivado",
                            ""
                        ))
                    conn.executemany("""
                        INSERT INTO transacoes (data, origem, descricao, categoria, subcategoria, tipo_movimentacao, valor_brl, responsavel, status, fitid)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """, registros_mig)
                    conn.commit()
                except Exception as err:
                    print(f"Erro na migracao inicial do CSV para SQLite: {err}")
                    semear_dados_demonstracao_sqlite()
            else:
                semear_dados_demonstracao_sqlite()

inicializar_banco()

def carregar_dados_sqlite() -> pd.DataFrame:
    with get_db_connection() as conn:
        df = pd.read_sql_query("SELECT * FROM transacoes ORDER BY id DESC", conn)
    return df

def inserir_transacao_sqlite(data_str, origem, descricao, categoria, subcategoria, tipo, valor, responsavel, fitid=""):
    with get_db_connection() as conn:
        conn.execute("""
            INSERT INTO transacoes (data, origem, descricao, categoria, subcategoria, tipo_movimentacao, valor_brl, responsavel, status, fitid)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'Efetivado', ?)
        """, (data_str, origem, sanitizar_texto(descricao), categoria, subcategoria, tipo, valor, responsavel, fitid))
        conn.commit()

# -------------------------------------------------------------
# 3. MOTOR DE REGRAS E CATEGORIZAÇÃO AUTOMÁTICA
# -------------------------------------------------------------
REGRAS = {
    "Transporte & Mobilidade": [r"movida", r"posto", r"combustivel", r"gasolina", r"uber", r"99app", r"estacionamento", r"brasil park"],
    "Alimentação - Padarias": [r"peter p[aã]o", r"celeiro do p[aã]o", r"padaria", r"caf[eé] com baguete", r"cacau show", r"confeitaria"],
    "Alimentação - Supermercados": [r"super mini", r"hortifruti", r"carrefour", r"atacadao", r"supermercado", r"fiorotti", r"extra"],
    "Alimentação - Restaurantes": [r"espetinho", r"sorveteria", r"restaurante", r"ifood", r"burger", r"rocon", r"lanchonete"],
    "Compras & Varejo": [r"mercado livre", r"americanas", r"tiktok shop", r"ponto sports", r"amazon", r"shopee", r"magalu"],
    "Assinaturas & Apps": [r"apple\.com", r"deezer", r"google play", r"meli\+", r"xbox", r"compra e volta", r"netflix", r"spotify", r"punko"],
    "Educação": [r"unintese", r"faculdade", r"curso", r"escola"],
    "Saúde & Farmácia": [r"drogasil", r"raia", r"drogaria", r"farmacia", r"pacheco"],
    "Tarifas Bancárias": [r"tarifa", r"iof", r"encargos", r"limite emergencial", r"anuidade"],
    "Família & Pessoal": [r"fabiola", r"fabíola", r"parente", r"familiar", r"dimas", r"repasse"],
    "Investimentos & Reserva": [r"cofrinho", r"reserva por gastos", r"dizimo", r"dízimo", r"cdi"]
}

def classificar_transacao(descricao: str, valor: float) -> tuple:
    desc = str(descricao).lower()
    
    if "transferencia mesma titularidade" in desc or "fluxo interno" in desc:
        return "Fluxo Interno", "Neutro"
    
    if "rendimento" in desc or "cdi" in desc:
        return "Receitas Externas", "Receita"
        
    for categoria, patterns in REGRAS.items():
        for pattern in patterns:
            if re.search(pattern, desc):
                if categoria == "Família & Pessoal":
                    tipo = "Receita" if valor > 0 else "Despesa"
                elif categoria == "Investimentos & Reserva":
                    tipo = "Reserva"
                else:
                    tipo = "Receita" if valor > 0 else "Despesa"
                return categoria, tipo
                
    tipo = "Receita" if valor > 0 else "Despesa"
    return "Outros / Diversos", tipo

def limpar_valor_monetario(val_raw) -> float:
    if isinstance(val_raw, (int, float)):
        return float(val_raw)
    s = str(val_raw).replace("R$", "").strip()
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    return float(s)

# -------------------------------------------------------------
# 4. PARSER NATIVO DE OFX (SGML / XML)
# -------------------------------------------------------------
def processar_conteudo_ofx(conteudo_texto: str) -> list:
    transacoes = []
    blocos = re.findall(
        r'<STMTTRN>(.*?)(?:</STMTTRN>|(?=<STMTTRN>)|(?=</BANKTRANLIST>)|$)', 
        conteudo_texto, 
        re.DOTALL | re.IGNORECASE
    )
    
    for bloco in blocos:
        m_type = re.search(r'<TRNTYPE>([^\r\n<]+)', bloco, re.IGNORECASE)
        trntype = m_type.group(1).strip().upper() if m_type else "OTHER"
        
        m_dt = re.search(r'<DTPOSTED>(\d{8})', bloco, re.IGNORECASE)
        if m_dt:
            dt_raw = m_dt.group(1)
            try:
                dt_obj = datetime.strptime(dt_raw, "%Y%m%d")
                data_str = dt_obj.strftime("%d/%m/%Y")
            except Exception:
                data_str = dt_raw
        else:
            data_str = date.today().strftime("%d/%m/%Y")
            
        m_val = re.search(r'<TRNAMT>([^\r\n<]+)', bloco, re.IGNORECASE)
        if m_val:
            val_str = m_val.group(1).strip().replace(",", ".")
            try:
                valor = float(val_str)
            except Exception:
                valor = 0.0
        else:
            valor = 0.0
            
        m_memo = re.search(r'<MEMO>([^\r\n<]+)', bloco, re.IGNORECASE)
        m_name = re.search(r'<NAME>([^\r\n<]+)', bloco, re.IGNORECASE)
        memo = m_memo.group(1).strip() if m_memo else ""
        name = m_name.group(1).strip() if m_name else ""
        descricao = memo or name or "Lançamento OFX"
        
        m_fitid = re.search(r'<FITID>([^\r\n<]+)', bloco, re.IGNORECASE)
        fitid = m_fitid.group(1).strip() if m_fitid else ""
        
        transacoes.append({
            "data": data_str,
            "descricao": sanitizar_texto(descricao),
            "valor": valor,
            "tipo_ofx": trntype,
            "fitid": fitid
        })
        
    return transacoes

def listar_arquivos_entrada() -> list:
    if not PASTA_ENTRADA.exists():
        return []
    extensoes = [".ofx", ".OFX", ".csv", ".CSV"]
    return [p for p in PASTA_ENTRADA.iterdir() if p.is_file() and p.suffix in extensoes]

def abrir_pasta_no_sistema(caminho: Path):
    try:
        if os.name == 'nt':
            os.startfile(str(caminho))
        elif sys.platform == 'darwin':
            subprocess.Popen(['open', str(caminho)])
        else:
            subprocess.Popen(['xdg-open', str(caminho)])
    except Exception as e:
        messagebox.showerror("Erro de Abertura", f"Não foi possível abrir o diretório:\n{e}")

# -------------------------------------------------------------
# 5. GERADOR AUTÔNOMO DE PLANILHA EXECUTIVA (OPENPYXL)
# -------------------------------------------------------------
def exportar_planilha_executiva_completa():
    df = carregar_dados_sqlite()
    if df.empty:
        messagebox.showwarning("Aviso", "Não há dados cadastrados no banco SQLite para exportar.")
        return None

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws_dash = wb.active
        ws_dash.title = "Dashboard Executivo"
        ws_trans = wb.create_sheet(title="Transações Conciliadas")

        # Paleta Corporativa Slate Dark & Ice Blue
        font_titulo = Font(name="Segoe UI", size=15, bold=True, color="0F172A")
        font_sub = Font(name="Segoe UI", size=10, italic=True, color="475569")
        font_hdr = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        font_bold = Font(name="Segoe UI", size=11, bold=True, color="0F172A")
        font_norm = Font(name="Segoe UI", size=10, color="1E293B")

        fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        fill_card = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

        thin_border = Border(
            left=Side(style='thin', color="CBD5E1"),
            right=Side(style='thin', color="CBD5E1"),
            top=Side(style='thin', color="CBD5E1"),
            bottom=Side(style='thin', color="CBD5E1")
        )

        # Aba 1: Dashboard Resumo
        ws_dash.views.sheetView[0].showGridLines = True
        ws_dash["B2"] = "💎 FinanceAgent - Relatório Executivo Consolidado"
        ws_dash["B2"].font = font_titulo
        ws_dash["B3"] = f"Exportação gerada em {datetime.now().strftime('%d/%m/%Y às %H:%M')} | Base: SQLite Nativo"
        ws_dash["B3"].font = font_sub

        total_rec = df[df["tipo_movimentacao"] == "Receita"]["valor_brl"].sum()
        total_desp = abs(df[df["tipo_movimentacao"] == "Despesa"]["valor_brl"].sum())
        total_res = abs(df[df["tipo_movimentacao"] == "Reserva"]["valor_brl"].sum())
        saldo_liq = total_rec - total_desp

        kpis = [
            ("ENTRADAS TOTAIS", total_rec, "B5", "C5"),
            ("DESPESAS TOTAIS", total_desp, "D5", "E5"),
            ("RESERVAS / APORTES", total_res, "F5", "G5"),
            ("SALDO LÍQUIDO", saldo_liq, "H5", "I5")
        ]

        for label, val, c1, c2 in kpis:
            ws_dash[c1] = label
            ws_dash[c1].font = Font(name="Segoe UI", size=9, bold=True, color="64748B")
            ws_dash[c2] = val
            ws_dash[c2].font = Font(name="Segoe UI", size=13, bold=True, color="0F172A" if label != "DESPESAS TOTAIS" else "DC2626")
            ws_dash[c2].number_format = '"R$" #,##0.00'
            ws_dash[c1].fill = fill_card
            ws_dash[c2].fill = fill_card
            ws_dash[c1].border = thin_border
            ws_dash[c2].border = thin_border

        # Resumo por Categoria
        ws_dash["B8"] = "RESUMO POR CATEGORIA"
        ws_dash["B8"].font = font_bold
        ws_dash.append([])

        cat_summary = df[df["tipo_movimentacao"] == "Despesa"].groupby("categoria")["valor_brl"].sum().abs().reset_index()
        cat_summary.columns = ["Categoria", "Total Despesa (R$)"]
        cat_summary = cat_summary.sort_values(by="Total Despesa (R$)", ascending=False)

        ws_dash.cell(row=9, column=2, value="Categoria").font = font_hdr
        ws_dash.cell(row=9, column=2).fill = fill_header
        ws_dash.cell(row=9, column=3, value="Total Gasto (R$)").font = font_hdr
        ws_dash.cell(row=9, column=3).fill = fill_header

        for idx, r in cat_summary.iterrows():
            r_idx = 10 + idx
            c_cat = ws_dash.cell(row=r_idx, column=2, value=r["Categoria"])
            c_val = ws_dash.cell(row=r_idx, column=3, value=r["Total Despesa (R$)"])
            c_cat.font = font_norm
            c_val.font = font_norm
            c_val.number_format = '"R$" #,##0.00'
            c_cat.border = thin_border
            c_val.border = thin_border
            if idx % 2 == 1:
                c_cat.fill = fill_zebra
                c_val.fill = fill_zebra

        # Aba 2: Transações Consolidadas
        ws_trans.views.sheetView[0].showGridLines = True
        headers_trans = ["ID", "Data", "Origem", "Descrição", "Categoria", "Subcategoria", "Tipo", "Valor (R$)", "Titular", "FITID"]
        ws_trans.append(headers_trans)

        for col_idx in range(1, len(headers_trans) + 1):
            cell = ws_trans.cell(row=1, column=col_idx)
            cell.font = font_hdr
            cell.fill = fill_header
            cell.alignment = Alignment(horizontal="center" if col_idx in [1, 2, 7, 9] else "left")

        for row_idx, r in df.iterrows():
            linha_excel = [
                r.get("id"),
                r.get("data"),
                r.get("origem"),
                r.get("descricao"),
                r.get("categoria"),
                r.get("subcategoria"),
                r.get("tipo_movimentacao"),
                r.get("valor_brl"),
                r.get("responsavel"),
                r.get("fitid")
            ]
            ws_trans.append(linha_excel)
            r_num = row_idx + 2
            
            c_val = ws_trans.cell(row=r_num, column=8)
            c_val.number_format = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
            
            for c_idx in range(1, len(headers_trans) + 1):
                c = ws_trans.cell(row=r_num, column=c_idx)
                c.font = font_norm
                c.border = thin_border
                if row_idx % 2 == 1:
                    c.fill = fill_zebra

        # Auto-ajuste de largura de colunas
        for ws in [ws_dash, ws_trans]:
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        nome_arquivo = f"FinanceAgent_Relatorio_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        caminho_final = PASTA_SAIDA / nome_arquivo
        wb.save(str(caminho_final))
        return caminho_final

    except Exception as err:
        messagebox.showerror("Erro na Exportação", f"Falha ao gerar planilha Excel:\n{err}")
        return None

# -------------------------------------------------------------
# 6. CONFIGURAÇÃO VISUAL E TEMA DO CUSTOMTKINTER
# -------------------------------------------------------------
ctk.set_appearance_mode("dark")
ctk.set_default_color_theme("blue")

try:
    plt.rcParams.update({
        'text.color': '#e2e8f0',
        'axes.labelcolor': '#94a3b8',
        'xtick.color': '#94a3b8',
        'ytick.color': '#94a3b8',
        'figure.facecolor': '#111622',
        'axes.facecolor': '#111622',
        'font.family': 'sans-serif'
    })
except Exception as err_theme:
    print(f"Aviso ao carregar tema padrão do CustomTkinter: {err_theme}")

class FinanceAgentExecutiveApp(ctk.CTk):
    def __init__(self):
        super().__init__()

        # Configurações da Janela Principal
        self.title("FinanceAgent | Executive Command Center")
        self.geometry("1240x780")
        self.minsize(1080, 680)
        self.configure(fg_color="#0b0f17")

        # Grid Principal: Sidebar (Col 0) + Workspace Dinâmico (Col 1)
        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(0, weight=1)

        self.fig_canvas = None
        self.fig_donut_canvas = None

        self.construir_sidebar()
        self.construir_workspace()

        # Inicializa na tela de Dashboard
        self.navegar_para("dashboard")
        self.atualizar_dados()

    # ---------------------------------------------------------
    # SIDEBAR EXECUTIVA (NAV-BAR LATERAL MODERNA)
    # ---------------------------------------------------------
    def construir_sidebar(self):
        self.sidebar = ctk.CTkFrame(self, width=250, corner_radius=0, fg_color="#111622", border_width=0)
        self.sidebar.grid(row=0, column=0, sticky="nsew")
        self.sidebar.grid_rowconfigure(10, weight=1)

        # Brand Container
        brand_frame = ctk.CTkFrame(self.sidebar, fg_color="transparent")
        brand_frame.grid(row=0, column=0, padx=20, pady=(24, 20), sticky="ew")

        lbl_logo = ctk.CTkLabel(
            brand_frame, 
            text="💎 FinanceAgent", 
            font=ctk.CTkFont(family="Segoe UI", size=17, weight="bold"),
            text_color="#38bdf8"
        )
        lbl_logo.pack(anchor="w")

        lbl_sub = ctk.CTkLabel(
            brand_frame, 
            text="Executive FinTech • Local-First", 
            font=ctk.CTkFont(family="Segoe UI", size=11),
            text_color="#64748b"
        )
        lbl_sub.pack(anchor="w", pady=(2, 0))

        # Divisor Sutil
        sep1 = ctk.CTkFrame(self.sidebar, height=1, fg_color="#1e293b")
        sep1.grid(row=1, column=0, padx=16, pady=(0, 16), sticky="ew")

        # Itens de Menu de Navegação
        self.btn_nav_dashboard = self._criar_botao_menu("📊  Dashboard Executivo", lambda: self.navegar_para("dashboard"), row=2)
        self.btn_nav_conciliar = self._criar_botao_menu("⚡  Conciliação OFX (Lote)", lambda: self.navegar_para("conciliacao"), row=3)
        self.btn_nav_extrato   = self._criar_botao_menu("📋  Extrato & Auditoria", lambda: self.navegar_para("extrato"), row=4)
        self.btn_nav_lancamento= self._criar_botao_menu("➕  Novo Lançamento", lambda: self.navegar_para("novo"), row=5)

        # Seção de Ações Rápidas de Pastas
        lbl_pastas = ctk.CTkLabel(
            self.sidebar, 
            text="DIRETÓRIOS LOCAIS", 
            font=ctk.CTkFont(size=9, weight="bold"),
            text_color="#475569"
        )
        lbl_pastas.grid(row=6, column=0, padx=20, pady=(24, 6), sticky="w")

        btn_pasta_ent = ctk.CTkButton(
            self.sidebar, 
            text="📂  Pasta Entrada (OFX/CSV)", 
            height=32, 
            fg_color="#1e293b", 
            hover_color="#334155",
            text_color="#94a3b8",
            anchor="w",
            font=ctk.CTkFont(size=11),
            command=lambda: abrir_pasta_no_sistema(PASTA_ENTRADA)
        )
        btn_pasta_ent.grid(row=7, column=0, padx=14, pady=3, sticky="ew")

        btn_pasta_sai = ctk.CTkButton(
            self.sidebar, 
            text="📂  Pasta Saída (Relatórios)", 
            height=32, 
            fg_color="#1e293b", 
            hover_color="#334155",
            text_color="#94a3b8",
            anchor="w",
            font=ctk.CTkFont(size=11),
            command=lambda: abrir_pasta_no_sistema(PASTA_SAIDA)
        )
        btn_pasta_sai.grid(row=8, column=0, padx=14, pady=3, sticky="ew")

        # Card de Status do Sistema no Rodapé
        self.status_card = ctk.CTkFrame(self.sidebar, fg_color="#0b0f17", corner_radius=10, border_width=1, border_color="#1e293b")
        self.status_card.grid(row=11, column=0, padx=14, pady=16, sticky="ew")

        status_header = ctk.CTkFrame(self.status_card, fg_color="transparent")
        status_header.pack(fill="x", padx=10, pady=(8, 2))

        ctk.CTkLabel(status_header, text="●", text_color="#10b981", font=ctk.CTkFont(size=13)).pack(side="left")
        ctk.CTkLabel(status_header, text=" SQLite Online", text_color="#e2e8f0", font=ctk.CTkFont(size=11, weight="bold")).pack(side="left", padx=4)

        self.lbl_status_desc = ctk.CTkLabel(
            self.status_card, 
            text="finanagent.db • Concorrente", 
            font=ctk.CTkFont(size=10), 
            text_color="#64748b"
        )
        self.lbl_status_desc.pack(anchor="w", padx=10, pady=(0, 8))

    def _criar_botao_menu(self, texto, comando, row):
        btn = ctk.CTkButton(
            self.sidebar, 
            text=texto, 
            height=38, 
            corner_radius=8,
            fg_color="transparent", 
            hover_color="#1e293b",
            text_color="#94a3b8",
            anchor="w",
            font=ctk.CTkFont(family="Segoe UI", size=12, weight="bold"),
            command=comando
        )
        btn.grid(row=row, column=0, padx=12, pady=3, sticky="ew")
        return btn

    # ---------------------------------------------------------
    # WORKSPACE PRINCIPAL E NAVEGAÇÃO DE VIEWS
    # ---------------------------------------------------------
    def construir_workspace(self):
        self.workspace = ctk.CTkFrame(self, fg_color="transparent")
        self.workspace.grid(row=0, column=1, sticky="nsew", padx=(18, 18), pady=18)
        self.workspace.grid_rowconfigure(1, weight=1)
        self.workspace.grid_columnconfigure(0, weight=1)

        # Header Superior do Workspace
        self.header_frame = ctk.CTkFrame(self.workspace, fg_color="#111622", corner_radius=12, height=60, border_width=1, border_color="#1e293b")
        self.header_frame.grid(row=0, column=0, sticky="ew", pady=(0, 14))

        self.lbl_view_title = ctk.CTkLabel(
            self.header_frame, 
            text="Dashboard Executivo & Análise Financeira", 
            font=ctk.CTkFont(family="Segoe UI", size=16, weight="bold"),
            text_color="#f8fafc"
        )
        self.lbl_view_title.pack(side="left", padx=18, pady=12)

        self.btn_atualizar_global = ctk.CTkButton(
            self.header_frame, 
            text="🔄  Recarregar", 
            width=110, 
            height=34,
            fg_color="#0284c7", 
            hover_color="#0369a1",
            font=ctk.CTkFont(size=11, weight="bold"),
            command=self.atualizar_dados
        )
        self.btn_atualizar_global.pack(side="right", padx=(6, 16), pady=12)

        self.btn_carregar_mock = ctk.CTkButton(
            self.header_frame, 
            text="✨  Carregar Dados Demo", 
            width=160, 
            height=34,
            fg_color="#10b981", 
            hover_color="#059669",
            font=ctk.CTkFont(size=11, weight="bold"),
            command=self.popular_dados_demo_ui
        )
        self.btn_carregar_mock.pack(side="right", padx=6, pady=12)

        self.btn_limpar_dados = ctk.CTkButton(
            self.header_frame, 
            text="🗑️  Limpar Base", 
            width=110, 
            height=34,
            fg_color="#334155", 
            hover_color="#475569",
            text_color="#cbd5e1",
            font=ctk.CTkFont(size=11),
            command=self.limpar_banco_ui
        )
        self.btn_limpar_dados.pack(side="right", padx=6, pady=12)

        # Container Central Dinâmico
        self.container_view = ctk.CTkFrame(self.workspace, fg_color="transparent")
        self.container_view.grid(row=1, column=0, sticky="nsew")
        self.container_view.grid_rowconfigure(0, weight=1)
        self.container_view.grid_columnconfigure(0, weight=1)

        # Inicializar as 4 Views Principais
        self.view_dashboard = ctk.CTkFrame(self.container_view, fg_color="transparent")
        self.view_conciliacao = ctk.CTkFrame(self.container_view, fg_color="transparent")
        self.view_extrato = ctk.CTkFrame(self.container_view, fg_color="transparent")
        self.view_novo = ctk.CTkFrame(self.container_view, fg_color="transparent")

        self._montar_view_dashboard()
        self._montar_view_conciliacao()
        self._montar_view_extrato()
        self._montar_view_novo()

    def navegar_para(self, tela_nome):
        # Reset visual dos botões
        botoes = [
            (self.btn_nav_dashboard, "dashboard"),
            (self.btn_nav_conciliar, "conciliacao"),
            (self.btn_nav_extrato, "extrato"),
            (self.btn_nav_lancamento, "novo")
        ]
        for btn, name in botoes:
            if name == tela_nome:
                btn.configure(fg_color="#1e293b", text_color="#38bdf8")
            else:
                btn.configure(fg_color="transparent", text_color="#94a3b8")

        # Esconder todas as views
        for v in [self.view_dashboard, self.view_conciliacao, self.view_extrato, self.view_novo]:
            v.grid_forget()

        # Exibir a selecionada
        if tela_nome == "dashboard":
            self.lbl_view_title.configure(text="Visão Executiva de Caixa & Inteligência Analítica")
            self.view_dashboard.grid(row=0, column=0, sticky="nsew")
            self.renderizar_graficos_matplotlib()
        elif tela_nome == "conciliacao":
            self.lbl_view_title.configure(text="Central de Conciliação em Lote (.OFX / .CSV)")
            self.view_conciliacao.grid(row=0, column=0, sticky="nsew")
            self.atualizar_lista_conciliacao()
        elif tela_nome == "extrato":
            self.lbl_view_title.configure(text="Extrato Consolidado & Auditoria de Transações")
            self.view_extrato.grid(row=0, column=0, sticky="nsew")
            self.recarregar_tabela_extrato()
        elif tela_nome == "novo":
            self.lbl_view_title.configure(text="Lançamento Manual Rápido (SQLite)")
            self.view_novo.grid(row=0, column=0, sticky="nsew")

    # ---------------------------------------------------------
    # VIEW 1: DASHBOARD EXECUTIVO COM KPIS E MATPLOTLIB NATIVO
    # ---------------------------------------------------------
    def _montar_view_dashboard(self):
        # 1. KPI CARDS (TOPO)
        kpi_frame = ctk.CTkFrame(self.view_dashboard, fg_color="transparent")
        kpi_frame.pack(fill="x", pady=(0, 14))
        kpi_frame.columnconfigure((0, 1, 2, 3), weight=1)

        self.kpi_rec, self.lbl_kpi_rec = self._criar_card_kpi(kpi_frame, "ENTRADAS TOTAIS", "R$ 0.00", "Receitas & Rendimentos", "#10b981", 0)
        self.kpi_desp, self.lbl_kpi_desp = self._criar_card_kpi(kpi_frame, "CONSUMO OPERACIONAL", "R$ 0.00", "Custo de vida do mês", "#f43f5e", 1)
        self.kpi_fam, self.lbl_kpi_fam = self._criar_card_kpi(kpi_frame, "FAMÍLIA & PESSOAL", "R$ 0.00", "Transferências / Apoio", "#38bdf8", 2)
        self.kpi_saldo, self.lbl_kpi_saldo = self._criar_card_kpi(kpi_frame, "RESULTADO DE CAIXA", "+R$ 0.00", "Saldo líquido consolidado", "#818cf8", 3)

        # 2. ÁREA DE GRÁFICOS ANALÍTICOS LADO A LADO
        charts_container = ctk.CTkFrame(self.view_dashboard, fg_color="transparent")
        charts_container.pack(fill="both", expand=True, pady=(0, 14))
        charts_container.columnconfigure(0, weight=3)
        charts_container.columnconfigure(1, weight=2)
        charts_container.rowconfigure(0, weight=1)

        # Frame Gráfico 1: Categorias (Barras Horizontais)
        self.frame_graf_cat = ctk.CTkFrame(charts_container, fg_color="#111622", corner_radius=12, border_width=1, border_color="#1e293b")
        self.frame_graf_cat.grid(row=0, column=0, sticky="nsew", padx=(0, 8))

        lbl_tit_cat = ctk.CTkLabel(
            self.frame_graf_cat, 
            text="📊  Distribuição de Gastos por Categoria (R$)", 
            font=ctk.CTkFont(size=12, weight="bold"),
            text_color="#f8fafc"
        )
        lbl_tit_cat.pack(anchor="w", padx=16, pady=(12, 6))

        # Frame Gráfico 2: Participação Titulares (Donut Chart)
        self.frame_graf_resp = ctk.CTkFrame(charts_container, fg_color="#111622", corner_radius=12, border_width=1, border_color="#1e293b")
        self.frame_graf_resp.grid(row=0, column=1, sticky="nsew", padx=(8, 0))

        lbl_tit_resp = ctk.CTkLabel(
            self.frame_graf_resp, 
            text="👥  Participação por Titular", 
            font=ctk.CTkFont(size=12, weight="bold"),
            text_color="#f8fafc"
        )
        lbl_tit_resp.pack(anchor="w", padx=16, pady=(12, 6))

        # 3. TERMINAL DE STATUS E LOGS EM TEMPO REAL
        term_frame = ctk.CTkFrame(self.view_dashboard, fg_color="#111622", corner_radius=12, border_width=1, border_color="#1e293b")
        term_frame.pack(fill="x", pady=(0, 4))

        self.txt_term_dash = ctk.CTkTextbox(
            term_frame, 
            height=68, 
            font=ctk.CTkFont(family="Consolas", size=10),
            fg_color="#090d13",
            text_color="#34d399",
            border_width=1,
            border_color="#1e293b"
        )
        self.txt_term_dash.pack(fill="both", expand=True, padx=12, pady=10)

    def _criar_card_kpi(self, parent, titulo, valor_padrao, subtitulo, cor_detalhe, col):
        card = ctk.CTkFrame(parent, fg_color="#111622", corner_radius=12, border_width=1, border_color="#1e293b")
        card.grid(row=0, column=col, padx=5, sticky="nsew")

        # Indicador de cor
        top_box = ctk.CTkFrame(card, fg_color="transparent")
        top_box.pack(fill="x", padx=16, pady=(14, 4))

        ctk.CTkLabel(top_box, text="■", text_color=cor_detalhe, font=ctk.CTkFont(size=12)).pack(side="left")
        ctk.CTkLabel(top_box, text=f" {titulo}", font=ctk.CTkFont(size=10, weight="bold"), text_color="#94a3b8").pack(side="left")

        lbl_valor = ctk.CTkLabel(
            card, 
            text=valor_padrao, 
            font=ctk.CTkFont(family="Segoe UI", size=20, weight="bold"),
            text_color="#f8fafc"
        )
        lbl_valor.pack(anchor="w", padx=16, pady=(0, 2))

        lbl_sub = ctk.CTkLabel(
            card, 
            text=subtitulo, 
            font=ctk.CTkFont(size=10),
            text_color="#64748b"
        )
        lbl_sub.pack(anchor="w", padx=16, pady=(0, 12))

        return card, lbl_valor

    def renderizar_graficos_matplotlib(self):
        # Limpar widgets anteriores
        for w in self.frame_graf_cat.winfo_children():
            if isinstance(w, tk.Widget) and not isinstance(w, ctk.CTkLabel):
                w.destroy()

        for w in self.frame_graf_resp.winfo_children():
            if isinstance(w, tk.Widget) and not isinstance(w, ctk.CTkLabel):
                w.destroy()

        df = carregar_dados_sqlite()
        if df.empty:
            ctk.CTkLabel(self.frame_graf_cat, text="Sem dados para gerar gráfico.", text_color="#64748b").pack(pady=40)
            ctk.CTkLabel(self.frame_graf_resp, text="Sem dados disponíveis.", text_color="#64748b").pack(pady=40)
            return

        df_desp = df[df["tipo_movimentacao"] == "Despesa"].copy()
        if df_desp.empty:
            ctk.CTkLabel(self.frame_graf_cat, text="Nenhuma despesa para análise.", text_color="#64748b").pack(pady=40)
            ctk.CTkLabel(self.frame_graf_resp, text="Nenhuma despesa para análise.", text_color="#64748b").pack(pady=40)
            return

        df_desp["valor_abs"] = df_desp["valor_brl"].abs()

        # 1. GRÁFICO DE BARRAS: DESPESAS POR CATEGORIA
        cat_agg = df_desp.groupby("categoria")["valor_abs"].sum().sort_values(ascending=True)

        fig1, ax1 = plt.subplots(figsize=(5.2, 3.0), dpi=100)
        fig1.patch.set_facecolor('#111622')
        ax1.set_facecolor('#111622')

        paleta_barras = ['#38bdf8', '#818cf8', '#34d399', '#f43f5e', '#fbbf24', '#a78bfa', '#f87171']
        cores = [paleta_barras[i % len(paleta_barras)] for i in range(len(cat_agg))]

        bars = ax1.barh(cat_agg.index, cat_agg.values, color=cores, height=0.6, edgecolor='none')

        # Customização Eixos
        ax1.tick_params(colors='#94a3b8', labelsize=8)
        ax1.spines['top'].set_visible(False)
        ax1.spines['right'].set_visible(False)
        ax1.spines['left'].set_color('#1e293b')
        ax1.spines['bottom'].set_color('#1e293b')
        ax1.grid(axis='x', color='#1e293b', linestyle='--', alpha=0.6)

        # Rótulos em R$ na ponta das barras
        for bar in bars:
            largura = bar.get_width()
            ax1.text(
                largura + (cat_agg.max() * 0.02),
                bar.get_y() + bar.get_height() / 2,
                f"R$ {largura:,.0f}",
                va='center', ha='left', color='#e2e8f0', fontsize=7.5, fontweight='bold'
            )

        ax1.set_xlim(0, cat_agg.max() * 1.22)
        fig1.tight_layout()

        canvas1 = FigureCanvasTkAgg(fig1, master=self.frame_graf_cat)
        canvas1.draw()
        canvas1.get_tk_widget().pack(fill="both", expand=True, padx=8, pady=(0, 6))

        # 2. GRÁFICO DONUT: DESPESAS POR TITULAR
        resp_agg = df_desp.groupby("responsavel")["valor_abs"].sum()

        fig2, ax2 = plt.subplots(figsize=(3.4, 3.0), dpi=100)
        fig2.patch.set_facecolor('#111622')
        ax2.set_facecolor('#111622')

        cores_resp = {'Usuário Titular': '#38bdf8', 'Usuário Secundário': '#f472b6', 'Compartilhado': '#34d399'}
        cores_donut = [cores_resp.get(k, '#818cf8') for k in resp_agg.index]

        wedges, texts, autotexts = ax2.pie(
            resp_agg.values, 
            labels=resp_agg.index, 
            autopct='%1.0f%%',
            startangle=140,
            colors=cores_donut,
            wedgeprops=dict(width=0.45, edgecolor='#111622', linewidth=2),
            pctdistance=0.75
        )

        for t in texts:
            t.set_color('#94a3b8')
            t.set_fontsize(8)
        for at in autotexts:
            at.set_color('#ffffff')
            at.set_fontsize(8)
            at.set_fontweight('bold')

        fig2.tight_layout()
        canvas2 = FigureCanvasTkAgg(fig2, master=self.frame_graf_resp)
        canvas2.draw()
        canvas2.get_tk_widget().pack(fill="both", expand=True, padx=8, pady=(0, 6))

    # ---------------------------------------------------------
    # VIEW 2: CENTRAL DE CONCILIAÇÃO OFX / CSV EM LOTE
    # ---------------------------------------------------------
    def _montar_view_conciliacao(self):
        panel = ctk.CTkFrame(self.view_conciliacao, fg_color="#111622", corner_radius=12, border_width=1, border_color="#1e293b")
        panel.pack(fill="both", expand=True, padx=4, pady=4)

        # Header de Ação com Botões de Destaque
        top_ctrl = ctk.CTkFrame(panel, fg_color="transparent")
        top_ctrl.pack(fill="x", padx=18, pady=(16, 12))

        ctk.CTkLabel(
            top_ctrl, 
            text="Fila de Extratos Prontos para Conciliação Automática", 
            font=ctk.CTkFont(size=14, weight="bold"),
            text_color="#f8fafc"
        ).pack(side="left")

        self.btn_executar_lote_main = ctk.CTkButton(
            top_ctrl, 
            text="🚀  Processar e Gravar Fila no SQLite", 
            height=38,
            fg_color="#10b981", 
            hover_color="#059669",
            text_color="#ffffff",
            font=ctk.CTkFont(size=12, weight="bold"),
            command=self.processar_arquivos_lote_ui
        )
        self.btn_executar_lote_main.pack(side="right", padx=(10, 0))

        self.btn_import_extratos_main = ctk.CTkButton(
            top_ctrl, 
            text="📂  Importar Extratos (.ofx / .csv)", 
            height=38,
            fg_color="#2563eb", 
            hover_color="#1d4ed8",
            text_color="#ffffff",
            font=ctk.CTkFont(size=12, weight="bold"),
            command=self.importar_e_copiar_extratos
        )
        self.btn_import_extratos_main.pack(side="right")

        # Seletores de Atribuição
        attr_frame = ctk.CTkFrame(panel, fg_color="#0b0f17", corner_radius=8)
        attr_frame.pack(fill="x", padx=18, pady=(0, 14))

        ctk.CTkLabel(attr_frame, text="Titular dos lançamentos:", text_color="#94a3b8", font=ctk.CTkFont(size=11)).pack(side="left", padx=(14, 8), pady=10)
        self.opt_titular_concil = ctk.CTkOptionMenu(attr_frame, values=["Usuário Titular", "Usuário Secundário", "Compartilhado"], fg_color="#1e293b", width=160)
        self.opt_titular_concil.pack(side="left", padx=(0, 16))
        self.opt_titular_concil.set("Usuário Titular")

        ctk.CTkLabel(attr_frame, text="Pasta Centralizada:", text_color="#94a3b8", font=ctk.CTkFont(size=11)).pack(side="left", padx=(0, 8))
        ctk.CTkLabel(attr_frame, text="entrada/ (Repositório Local)", text_color="#38bdf8", font=ctk.CTkFont(size=11, weight="bold")).pack(side="left")

        # Visualizador de Arquivos na Fila
        self.txt_lista_lote_main = ctk.CTkTextbox(
            panel, 
            font=ctk.CTkFont(family="Consolas", size=11),
            fg_color="#090d13",
            text_color="#38bdf8",
            border_width=1,
            border_color="#1e293b"
        )
        self.txt_lista_lote_main.pack(fill="both", expand=True, padx=18, pady=(0, 18))

    def atualizar_lista_conciliacao(self):
        arquivos = listar_arquivos_entrada()
        self.txt_lista_lote_main.delete("1.0", "end")

        if arquivos:
            msg = f"🔍 {len(arquivos)} arquivo(s) detectado(s) na pasta 'entrada/':\n"
            msg += "=" * 80 + "\n"
            for a in arquivos:
                tam_kb = a.stat().st_size / 1024
                dt_mod = datetime.fromtimestamp(a.stat().st_mtime).strftime("%d/%m/%Y %H:%M")
                msg += f" • [{a.suffix.upper()}]  {a.name:<45} | {tam_kb:>7.1f} KB | Modificado: {dt_mod}\n"
            msg += "=" * 80 + "\n"
            msg += "Clique em 'Processar e Gravar Fila no SQLite' para categorizar e conciliar automaticamente."
            self.txt_lista_lote_main.insert("1.0", msg)
        else:
            self.txt_lista_lote_main.insert("1.0", "Nenhum arquivo .ofx ou .csv encontrado na pasta 'entrada/'.\n\nDeposite seus extratos bancários na pasta 'entrada/' ou clique em 'Importar Arquivo Avulso'.")

    # ---------------------------------------------------------
    # VIEW 3: EXTRATO & AUDITORIA DE TRANSAÇÕES (TABELA TTK)
    # ---------------------------------------------------------
    def _montar_view_extrato(self):
        panel = ctk.CTkFrame(self.view_extrato, fg_color="#111622", corner_radius=12, border_width=1, border_color="#1e293b")
        panel.pack(fill="both", expand=True, padx=4, pady=4)

        # Barra de Filtros e Exportação
        bar = ctk.CTkFrame(panel, fg_color="transparent")
        bar.pack(fill="x", padx=16, pady=(14, 10))

        ctk.CTkLabel(bar, text="Filtro por Titular:", font=ctk.CTkFont(size=11), text_color="#94a3b8").pack(side="left", padx=(0, 6))
        self.opt_filtro_resp_main = ctk.CTkOptionMenu(
            bar, 
            values=["Todos", "Usuário Titular", "Usuário Secundário", "Compartilhado"], 
            width=150, 
            fg_color="#1e293b",
            command=lambda _: self.recarregar_tabela_extrato()
        )
        self.opt_filtro_resp_main.pack(side="left", padx=(0, 16))
        self.opt_filtro_resp_main.set("Todos")

        ctk.CTkLabel(bar, text="Buscar Transação:", font=ctk.CTkFont(size=11), text_color="#94a3b8").pack(side="left", padx=(0, 6))
        self.ent_busca_main = ctk.CTkEntry(bar, placeholder_text="Filtrar por descrição...", width=200, fg_color="#0b0f17")
        self.ent_busca_main.pack(side="left", padx=(0, 10))
        self.ent_busca_main.bind("<KeyRelease>", lambda e: self.recarregar_tabela_extrato())

        btn_limpar_busca = ctk.CTkButton(
            bar, 
            text="Limpar", 
            width=70, 
            fg_color="#334155", 
            hover_color="#475569",
            command=lambda: [self.ent_busca_main.delete(0, "end"), self.recarregar_tabela_extrato()]
        )
        btn_limpar_busca.pack(side="left")

        # Botão Exportar Excel Profissional
        btn_export_excel = ctk.CTkButton(
            bar, 
            text="📊  Exportar Relatório Excel (.xlsx)", 
            height=36,
            fg_color="#059669", 
            hover_color="#047857",
            font=ctk.CTkFont(size=11, weight="bold"),
            command=self.acao_exportar_excel
        )
        btn_export_excel.pack(side="right")

        # Tabela TTK Estilizada Dark
        tbl_frame = ctk.CTkFrame(panel, fg_color="#0b0f17", corner_radius=10, border_width=1, border_color="#1e293b")
        tbl_frame.pack(fill="both", expand=True, padx=16, pady=(0, 16))

        style = ttk.Style()
        style.theme_use("clam")
        style.configure("Treeview",
            background="#0f172a",
            foreground="#f8fafc",
            rowheight=28,
            fieldbackground="#0f172a",
            bordercolor="#1e293b",
            borderwidth=0,
            font=("Segoe UI", 10)
        )
        style.configure("Treeview.Heading",
            background="#1e293b",
            foreground="#38bdf8",
            relief="flat",
            font=("Segoe UI", 10, "bold")
        )
        style.map("Treeview",
            background=[('selected', '#2563eb')],
            foreground=[('selected', '#ffffff')]
        )

        colunas = ("id", "data", "origem", "descricao", "categoria", "tipo", "valor", "responsavel")
        self.tree_main = ttk.Treeview(tbl_frame, columns=colunas, show="headings", selectmode="browse")

        self.tree_main.heading("id", text="ID")
        self.tree_main.heading("data", text="Data")
        self.tree_main.heading("origem", text="Origem / Conta")
        self.tree_main.heading("descricao", text="Descrição do Lançamento")
        self.tree_main.heading("categoria", text="Categoria Analítica")
        self.tree_main.heading("tipo", text="Tipo")
        self.tree_main.heading("valor", text="Valor (R$)")
        self.tree_main.heading("responsavel", text="Titular")

        self.tree_main.column("id", width=45, anchor="center")
        self.tree_main.column("data", width=85, anchor="center")
        self.tree_main.column("origem", width=120, anchor="w")
        self.tree_main.column("descricao", width=250, anchor="w")
        self.tree_main.column("categoria", width=170, anchor="w")
        self.tree_main.column("tipo", width=80, anchor="center")
        self.tree_main.column("valor", width=100, anchor="e")
        self.tree_main.column("responsavel", width=120, anchor="center")

        sb = ttk.Scrollbar(tbl_frame, orient="vertical", command=self.tree_main.yview)
        self.tree_main.configure(yscrollcommand=sb.set)

        self.tree_main.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")

    def recarregar_tabela_extrato(self):
        for item in self.tree_main.get_children():
            self.tree_main.delete(item)

        df = carregar_dados_sqlite()
        if df.empty:
            return

        filtro = self.opt_filtro_resp_main.get()
        if filtro != "Todos":
            df = df[df["responsavel"] == filtro]

        busca = self.ent_busca_main.get().strip().lower()
        if busca:
            df = df[df["descricao"].str.lower().str.contains(busca, na=False)]

        for _, row in df.iterrows():
            v_num = float(row.get("valor_brl", 0.0))
            v_formatado = f"R$ {v_num:,.2f}"
            self.tree_main.insert("", "end", values=(
                row.get("id"),
                row.get("data"),
                row.get("origem"),
                row.get("descricao"),
                row.get("categoria"),
                row.get("tipo_movimentacao"),
                v_formatado,
                row.get("responsavel")
            ))

    # ---------------------------------------------------------
    # VIEW 4: NOVO LANÇAMENTO MANUAL (CLEAN FORM)
    # ---------------------------------------------------------
    def _montar_view_novo(self):
        panel = ctk.CTkFrame(self.view_novo, fg_color="#111622", corner_radius=12, border_width=1, border_color="#1e293b")
        panel.pack(fill="both", expand=True, padx=4, pady=4)

        form_box = ctk.CTkFrame(panel, fg_color="#0b0f17", corner_radius=10, border_width=1, border_color="#1e293b")
        form_box.pack(padx=40, pady=30, fill="both", expand=True)

        ctk.CTkLabel(
            form_box, 
            text="Inserir Registro Manual no SQLite", 
            font=ctk.CTkFont(size=16, weight="bold"),
            text_color="#f8fafc"
        ).pack(anchor="w", padx=24, pady=(20, 16))

        grid = ctk.CTkFrame(form_box, fg_color="transparent")
        grid.pack(fill="x", padx=24, pady=10)
        grid.columnconfigure((0, 1), weight=1)

        # Data e Titular
        ctk.CTkLabel(grid, text="Data do Gasto:", text_color="#94a3b8").grid(row=0, column=0, sticky="w", pady=(4, 2))
        self.ent_novo_data = ctk.CTkEntry(grid, fg_color="#161b22")
        self.ent_novo_data.insert(0, date.today().strftime("%d/%m/%Y"))
        self.ent_novo_data.grid(row=1, column=0, sticky="ew", padx=(0, 16), pady=(0, 14))

        ctk.CTkLabel(grid, text="Responsável / Titular:", text_color="#94a3b8").grid(row=0, column=1, sticky="w", pady=(4, 2))
        self.opt_novo_resp = ctk.CTkOptionMenu(grid, values=["Usuário Titular", "Usuário Secundário", "Compartilhado"], fg_color="#161b22")
        self.opt_novo_resp.grid(row=1, column=1, sticky="ew", pady=(0, 14))
        self.opt_novo_resp.set("Usuário Titular")

        # Descrição
        ctk.CTkLabel(grid, text="Descrição / Estabelecimento:", text_color="#94a3b8").grid(row=2, column=0, columnspan=2, sticky="w", pady=(4, 2))
        self.ent_novo_desc = ctk.CTkEntry(grid, placeholder_text="Ex: Supermercado Carrefour, Farmácia Raia...", fg_color="#161b22")
        self.ent_novo_desc.grid(row=3, column=0, columnspan=2, sticky="ew", pady=(0, 14))

        # Tipo e Valor
        ctk.CTkLabel(grid, text="Tipo de Movimento:", text_color="#94a3b8").grid(row=4, column=0, sticky="w", pady=(4, 2))
        self.opt_novo_tipo = ctk.CTkOptionMenu(grid, values=["Despesa", "Receita", "Neutro"], fg_color="#161b22")
        self.opt_novo_tipo.grid(row=5, column=0, sticky="ew", padx=(0, 16), pady=(0, 14))
        self.opt_novo_tipo.set("Despesa")

        ctk.CTkLabel(grid, text="Valor em R$:", text_color="#94a3b8").grid(row=4, column=1, sticky="w", pady=(4, 2))
        self.ent_novo_val = ctk.CTkEntry(grid, placeholder_text="Ex: 150,00", fg_color="#161b22")
        self.ent_novo_val.grid(row=5, column=1, sticky="ew", pady=(0, 14))

        # Forma de Pagamento
        ctk.CTkLabel(grid, text="Forma de Pagamento / Origem:", text_color="#94a3b8").grid(row=6, column=0, sticky="w", pady=(4, 2))
        self.opt_novo_origem = ctk.CTkOptionMenu(grid, values=["Cartão de Crédito", "Conta Débito / Pix", "Dinheiro"], fg_color="#161b22")
        self.opt_novo_origem.grid(row=7, column=0, sticky="ew", padx=(0, 16), pady=(0, 24))
        self.opt_novo_origem.set("Cartão de Crédito")

        # Botão Salvar
        btn_salvar = ctk.CTkButton(
            form_box, 
            text="💾  Gravar Lançamento no Banco SQLite", 
            height=42,
            fg_color="#10b981", 
            hover_color="#059669",
            font=ctk.CTkFont(size=13, weight="bold"),
            command=self.salvar_lancamento_manual
        )
        btn_salvar.pack(padx=24, pady=(0, 24), fill="x")

    # ---------------------------------------------------------
    # OPERAÇÕES DE DADOS E EVENTOS
    # ---------------------------------------------------------
    def atualizar_dados(self):
        df = carregar_dados_sqlite()
        arquivos_fila = listar_arquivos_entrada()

        # Atualizar KPIs
        if not df.empty:
            rec = df[df["tipo_movimentacao"] == "Receita"]["valor_brl"].sum()
            desp_op = abs(df[
                (df["tipo_movimentacao"] == "Despesa") & 
                (~df["categoria"].str.contains("Família", na=False))
            ]["valor_brl"].sum())
            fam = abs(df[df["categoria"].str.contains("Família", na=False)]["valor_brl"].sum())
            saldo = rec - desp_op - fam
        else:
            rec, desp_op, fam, saldo = 0.0, 0.0, 0.0, 0.0

        self.lbl_kpi_rec.configure(text=f"R$ {rec:,.2f}")
        self.lbl_kpi_desp.configure(text=f"R$ {desp_op:,.2f}")
        self.lbl_kpi_fam.configure(text=f"R$ {fam:,.2f}")
        sinal = "+" if saldo >= 0 else ""
        self.lbl_kpi_saldo.configure(text=f"{sinal}R$ {saldo:,.2f}")

        # Atualizar Terminal de Atividades
        total_ofx = len([f for f in arquivos_fila if f.suffix.lower() == ".ofx"])
        total_csv = len([f for f in arquivos_fila if f.suffix.lower() == ".csv"])

        texto_terminal = f"""[SISTEMA CONECTADO] Sessão de Auditoria Financeira Ativa
• Base Local: {DB_PATH.name} ({len(df)} registros ativos) | Pasta Entrada: {total_ofx} .OFX e {total_csv} .CSV na fila.
• Diretório de Trabalho: {BASE_DIR}"""
        self.txt_term_dash.delete("1.0", "end")
        self.txt_term_dash.insert("1.0", texto_terminal)

        self.renderizar_graficos_matplotlib()
        self.atualizar_lista_conciliacao()
        self.recarregar_tabela_extrato()

    def popular_dados_demo_ui(self):
        semear_dados_demonstracao_sqlite()
        self.atualizar_dados()
        messagebox.showinfo("Demonstração Carregada", "✅ Dados fictícios de demonstração inseridos com sucesso!\nO Dashboard, gráficos e tabelas foram populados.")

    def limpar_banco_ui(self):
        resp = messagebox.askyesno("Confirmar Limpeza", "Deseja realmente apagar todos os registros do banco local?\nEsta ação deixará a base limpa para novas importações.")
        if resp:
            with get_db_connection() as conn:
                conn.execute("DELETE FROM transacoes")
                conn.commit()
            self.atualizar_dados()
            messagebox.showinfo("Base Limpa", "🗑️ Todos os registros foram removidos com sucesso.")

    def salvar_lancamento_manual(self):
        desc = self.ent_novo_desc.get().strip()
        val_str = self.ent_novo_val.get().strip()
        data_str = self.ent_novo_data.get().strip()
        resp = self.opt_novo_resp.get()
        tipo = self.opt_novo_tipo.get()
        origem = self.opt_novo_origem.get()

        if not desc:
            messagebox.showwarning("Aviso", "Informe a descrição do lançamento.")
            return

        try:
            val_num = limpar_valor_monetario(val_str)
        except Exception:
            messagebox.showerror("Erro", "Valor financeiro inválido.")
            return

        val_final = -abs(val_num) if tipo == "Despesa" else abs(val_num)
        cat_calc, _ = classificar_transacao(desc, val_final)

        inserir_transacao_sqlite(
            data_str=data_str,
            origem=origem,
            descricao=desc,
            categoria=cat_calc,
            subcategoria="",
            tipo=tipo,
            valor=val_final,
            responsavel=resp
        )

        messagebox.showinfo("Sucesso", "✅ Lançamento gravado no SQLite com sucesso!")
        self.ent_novo_desc.delete(0, "end")
        self.ent_novo_val.delete(0, "end")
        self.atualizar_dados()
        self.navegar_para("extrato")

    def processar_arquivos_lote_ui(self):
        arquivos = listar_arquivos_entrada()
        if not arquivos:
            messagebox.showwarning("Fila Vazia", "Nenhum arquivo .ofx ou .csv encontrado na pasta 'entrada/'.")
            return

        titular_selecionado = self.opt_titular_concil.get()
        total_importados = 0
        arquivos_ok = 0
        novos_registros = []

        for arq in arquivos:
            try:
                if arq.suffix.lower() == ".ofx":
                    conteudo = arq.read_text(encoding="utf-8", errors="ignore")
                    trans_ofx = processar_conteudo_ofx(conteudo)
                    for item in trans_ofx:
                        cat_calc, tipo_calc = classificar_transacao(item["descricao"], item["valor"])
                        novos_registros.append((
                            item["data"],
                            f"OFX ({arq.name})",
                            item["descricao"],
                            cat_calc,
                            "",
                            tipo_calc,
                            item["valor"],
                            titular_selecionado,
                            "Efetivado",
                            item["fitid"]
                        ))
                    total_importados += len(trans_ofx)
                    arquivos_ok += 1
                elif arq.suffix.lower() == ".csv":
                    df_c = pd.read_csv(arq)
                    col_desc = next((c for c in df_c.columns if c.lower() in ["descricao", "estabelecimento", "historico", "lancamento"]), None)
                    col_val = next((c for c in df_c.columns if c.lower() in ["valor_brl", "valor", "montante", "quantia"]), None)
                    col_data = next((c for c in df_c.columns if c.lower() in ["data", "dt"]), None)
                    col_resp = next((c for c in df_c.columns if c.lower() in ["responsavel", "titular"]), None)
                    
                    if col_desc and col_val:
                        for _, r in df_c.iterrows():
                            desc_val = str(r[col_desc])
                            val_num = limpar_valor_monetario(r[col_val])
                            dt_str = str(r[col_data]) if col_data else date.today().strftime("%d/%m/%Y")
                            resp_val = str(r[col_resp]) if col_resp else titular_selecionado
                            
                            cat_calc, tipo_calc = classificar_transacao(desc_val, val_num)
                            novos_registros.append((
                                dt_str,
                                f"CSV ({arq.name})",
                                sanitizar_texto(desc_val),
                                cat_calc,
                                "",
                                tipo_calc,
                                val_num,
                                resp_val,
                                "Efetivado",
                                ""
                            ))
                        total_importados += len(df_c)
                        arquivos_ok += 1
            except Exception as e:
                print(f"Erro ao processar {arq.name}: {e}")

        if novos_registros:
            with get_db_connection() as conn:
                conn.executemany("""
                    INSERT INTO transacoes (data, origem, descricao, categoria, subcategoria, tipo_movimentacao, valor_brl, responsavel, status, fitid)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, novos_registros)
                conn.commit()

            messagebox.showinfo(
                "Conciliação em Lote Concluída", 
                f"✅ Sucesso!\n• Arquivos processados: {arquivos_ok}\n• Transações gravadas: {total_importados}\n• Titular atribuído: {titular_selecionado}"
            )
            self.atualizar_dados()
            self.navegar_para("dashboard")

    def importar_e_copiar_extratos(self):
        caminhos = filedialog.askopenfilenames(
            title="Selecione um ou mais Extratos Bancários (.OFX / .CSV)",
            filetypes=[("Extratos Bancários", "*.ofx *.csv *.OFX *.CSV"), ("Arquivos OFX", "*.ofx *.OFX"), ("Arquivos CSV", "*.csv *.CSV")]
        )
        if not caminhos:
            return

        total_copiados = 0
        for c in caminhos:
            p = Path(c)
            destino = PASTA_ENTRADA / p.name
            try:
                shutil.copy2(p, destino)
                total_copiados += 1
            except Exception as err:
                print(f"Erro ao copiar {p.name}: {err}")

        messagebox.showinfo("Importação Concluída", f"📁 {total_copiados} arquivo(s) copiado(s) para a pasta 'entrada/'.")
        self.atualizar_dados()
        self.atualizar_lista_conciliacao()

    def acao_exportar_excel(self):
        caminho = exportar_planilha_executiva_completa()
        if caminho and caminho.exists():
            resp = messagebox.askyesno(
                "Exportação Concluída com Sucesso",
                f"📊 Relatório Excel gerado:\n{caminho.name}\n\nLocal:\n{PASTA_SAIDA}\n\nDeseja abrir o arquivo agora?"
            )
            if resp:
                abrir_pasta_no_sistema(caminho)

# -------------------------------------------------------------
# 7. EXECUÇÃO PRINCIPAL
# -------------------------------------------------------------
if __name__ == "__main__":
    app = FinanceAgentExecutiveApp()
    app.mainloop()
