# -*- coding: utf-8 -*-
"""
FinanceAgent - Gerador de Planilha Excel com Dashboard e Base Consolidada
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import PieChart, Reference

def gerar_planilha():
    wb = openpyxl.Workbook()

    # Setup sheets
    ws_dash = wb.active
    ws_dash.title = "Dashboard Mensal"
    ws_trans = wb.create_sheet(title="Transações Consolidadas")

    font_title = Font(name="Calibri", size=15, bold=True, color="1F4E79")
    font_subtitle = Font(name="Calibri", size=11, italic=True, color="595959")
    font_tbl_hdr = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True, color="000000")
    font_regular = Font(name="Calibri", size=11, color="000000")

    fill_steel = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    fill_light_blue = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    fill_zebra = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    fill_card_bg = PatternFill(start_color="F2F4F8", end_color="F2F4F8", fill_type="solid")
    fill_green_soft = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    fill_amber_soft = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    thin_gray = Side(style='thin', color="D9D9D9")
    border_cell = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    # Base de Dados
    trans_headers = ["Data", "Origem / Conta", "Estabelecimento / Descrição", "Categoria Macro", "Subcategoria", "Tipo Movimentação", "Valor (R$)"]
    ws_trans.append([])
    ws_trans.append(["FinanceAgent - Base Consolidada de Movimentações"])
    ws_trans.merge_cells("A2:G2")
    ws_trans["A2"].font = font_title

    ws_trans.append(trans_headers)
    for col_num in range(1, len(trans_headers) + 1):
        c = ws_trans.cell(row=3, column=col_num)
        c.font = font_tbl_hdr
        c.fill = fill_steel
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border_cell

    raw_data = [
        ("01/08/2026", "Conta Bancária", "Brasil Park Estacionamento", "Transporte & Mobilidade", "Estacionamento", "Despesa", -10.00),
        ("02/08/2026", "Conta Bancária", "Prestador de Servicos Gerais", "Serviços & Terceiros", "Prestador de Serviço", "Despesa", -82.00),
        ("03/08/2026", "Conta Bancária", "Recebimento Servicos Terceiros", "Receitas Externas", "Terceiros", "Receita", 125.00),
        ("04/08/2026", "Conta Bancária", "Recebimento Venda Produto", "Receitas Externas", "Comércio", "Receita", 26.00),
        ("05/08/2026", "Conta Bancária", "Associação Educacional EAD", "Educação", "Cursos / Faculdade", "Despesa", -130.05),
        ("07/08/2026", "Conta Bancária", "Aporte Mesma Titularidade", "Fluxo Interno", "Aporte Próprio", "Neutro", 204.38),
        ("08/08/2026", "Conta Bancária", "Transferencia Mesma Titularidade", "Fluxo Interno", "Transferência Própria", "Neutro", -320.00),
        ("13/08/2026", "Conta Bancária", "Recebimento Consultoria", "Receitas Externas", "Serviços", "Receita", 30.00),
        ("13/08/2026", "Conta Bancária", "Serviço Manutenção Técnica", "Serviços & Terceiros", "Prestador de Serviço", "Despesa", -5.00),
        ("14/08/2026", "Conta Bancária", "Aporte Mesma Titularidade", "Fluxo Interno", "Aporte Próprio", "Neutro", 1400.00),
        ("14/08/2026", "Conta Bancária", "Repasse Familiar Mensal", "Família & Pessoal", "Ajuda Familiar", "Despesa", -1500.00),
        ("15/07/2026", "Cartão Visa", "Raia Drogasil Farmácia", "Saúde & Farmácia", "Farmácia", "Crédito", -20.07),
        ("22/07/2026", "Cartão Visa", "Posto Combustível Shell", "Transporte & Mobilidade", "Combustível", "Crédito", -100.00),
        ("22/07/2026", "Cartão Visa", "Lojas Americanas Varejo", "Compras & Varejo", "Lojas Físicas", "Crédito", -62.69),
        ("22/07/2026", "Cartão Visa", "Movida Carro Assinatura", "Transporte & Mobilidade", "Carro Assinatura", "Crédito", -1924.12),
        ("22/07/2026", "Cartão Visa", "Mercado Livre Marketplace", "Compras & Varejo", "E-commerce", "Crédito", -82.61),
        ("23/07/2026", "Cartão Visa", "Padaria Peter Pão", "Alimentação", "Padarias & Docerias", "Crédito", -23.46),
        ("24/07/2026", "Cartão Visa", "Restaurante Espetinho Gourmet", "Alimentação", "Restaurantes & Lanches", "Crédito", -35.00),
        ("24/07/2026", "Cartão Visa", "Padaria Peter Pão", "Alimentação", "Padarias & Docerias", "Crédito", -13.19),
        ("25/07/2026", "Cartão Visa", "Supermercado Super Mini", "Alimentação", "Supermercados", "Crédito", -23.97),
        ("26/07/2026", "Cartão Visa", "Padaria Peter Pão", "Alimentação", "Padarias & Docerias", "Crédito", -13.79),
        ("26/07/2026", "Cartão Visa", "Microsoft Xbox Assinatura", "Lazer & Assinaturas", "Jogos & Apps", "Crédito", -37.48),
        ("28/07/2026", "Cartão Visa", "Cacau Show Chocolates", "Alimentação", "Padarias & Docerias", "Crédito", -51.23),
        ("29/07/2026", "Cartão Visa", "Padaria Peter Pão", "Alimentação", "Padarias & Docerias", "Crédito", -5.53),
        ("02/08/2026", "Cartão Visa", "Google Play Assinaturas", "Assinaturas & Apps", "Serviços Digitais", "Crédito", -9.90),
        ("03/08/2026", "Cartão Visa", "Posto Combustível Ipiranga", "Transporte & Mobilidade", "Combustível", "Crédito", -100.00),
        ("04/08/2026", "Cartão Visa", "Deezer Music Streaming", "Assinaturas & Apps", "Streaming", "Crédito", -24.90),
        ("07/08/2026", "Cartão Visa", "Padaria Celeiro do Pão", "Alimentação", "Padarias & Docerias", "Crédito", -40.43),
        ("08/08/2026", "Cartão Visa", "Café com Baguete Doceria", "Alimentação", "Padarias & Docerias", "Crédito", -38.00),
        ("08/08/2026", "Cartão Visa", "Posto Combustível BR", "Transporte & Mobilidade", "Combustível", "Crédito", -100.00),
        ("10/08/2026", "Cartão Visa", "Ponto Sports Vestuário", "Compras & Varejo", "Vestuário", "Crédito", -109.90),
        ("12/08/2026", "Cartão Visa", "Google Play Aplicativos", "Assinaturas & Apps", "Jogos & Apps", "Crédito", -6.98),
        ("13/08/2026", "Cartão Visa", "Supermercado Super Mini", "Alimentação", "Supermercados", "Crédito", -9.69),
        ("13/08/2026", "Cartão Visa", "Hortifruti Fiorotti", "Alimentação", "Supermercados", "Crédito", -8.98),
        ("16/08/2026", "Conta Bancária", "Aporte Mesma Titularidade", "Fluxo Interno", "Aporte Próprio", "Neutro", 1621.00),
        ("16/08/2026", "Fatura Crédito", "Tarifa Limite Emergencial", "Tarifas Bancárias", "Encargos Cartão", "Crédito", -14.90),
        ("17/08/2026", "Cartão Visa", "Mercado Livre E-commerce", "Compras & Varejo", "E-commerce", "Crédito", -37.28),
        ("17/08/2026", "Cartão Visa", "Compra e Volta Cashback", "Assinaturas & Apps", "Serviços Digitais", "Crédito", -27.00),
        ("19/08/2026", "Conta Bancária", "Recebimento Repasse Parceiro", "Receitas Externas", "Parcerias", "Receita", 1500.00),
        ("20/08/2026", "Conta Bancária", "Padaria Peter Pão", "Alimentação", "Padarias & Docerias", "Despesa", -15.36),
        ("21/08/2026", "Cartão Visa", "Meli+ Assinatura Mensal", "Assinaturas & Apps", "Serviços Digitais", "Crédito", -9.90),
        ("21/08/2026", "Conta Bancária", "Sorveteria Rocon", "Alimentação", "Restaurantes & Lanches", "Despesa", -7.00),
        ("22/08/2026", "Conta Bancária", "Padaria Peter Pão", "Alimentação", "Padarias & Docerias", "Despesa", -22.52),
        ("28/08/2026", "Conta Bancária", "Recebimento Repasse Parceiro", "Receitas Externas", "Parcerias", "Receita", 200.00),
        ("28/08/2026", "Conta Bancária", "Serviço Manutenção Externa", "Serviços & Terceiros", "Prestador de Serviço", "Despesa", -90.00),
        ("29/08/2026", "Conta Bancária", "Supermercado Super Mini", "Alimentação", "Supermercados", "Despesa", -9.98),
        ("30/08/2026", "Conta Bancária", "Aporte Mesma Titularidade", "Fluxo Interno", "Aporte Próprio", "Neutro", 189.32),
        ("31/08/2026", "Conta Bancária", "Reserva Automática Cofrinho", "Investimentos & Reserva", "Dízimo / Cofrinho", "Reserva", -34.00),
        ("31/08/2026", "Conta Bancária", "Rendimento CDI Saldo", "Receitas Externas", "Rendimento Saldo", "Receita", 2.27),
    ]

    for row_idx, r in enumerate(raw_data, start=4):
        ws_trans.append(r)
        val_cell = ws_trans.cell(row=row_idx, column=7)
        val_cell.number_format = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
        for c_idx in range(1, 8):
            cell = ws_trans.cell(row=row_idx, column=c_idx)
            cell.border = border_cell
            cell.font = font_regular
            if row_idx % 2 == 0:
                cell.fill = fill_zebra

    # Ajuste colunas da aba transacoes
    for col in ws_trans.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_trans.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # ABA DASHBOARD
    ws_dash.views.sheetView[0].showGridLines = True
    ws_dash.append([])
    ws_dash.append(["", "PAINEL EXECUTIVO DE CONTROLE FINANCEIRO"])
    ws_dash["B2"].font = font_title
    ws_dash.append(["", "Visão Analítica Consolidada"])
    ws_dash["B3"].font = font_subtitle

    # Cards KPI
    cards = [
        ("ENTRADAS TOTAIS", '=SUMIFS(\'Transações Consolidadas\'!G:G, \'Transações Consolidadas\'!F:F, "Receita")', "B5", "C5"),
        ("CONSUMO OPERACIONAL", '=ABS(SUMIFS(\'Transações Consolidadas\'!G:G, \'Transações Consolidadas\'!F:F, "<>Receita", \'Transações Consolidadas\'!F:F, "<>Neutro", \'Transações Consolidadas\'!D:D, "<>Família & Pessoal", \'Transações Consolidadas\'!F:F, "<>Reserva"))', "D5", "E5"),
        ("FAMÍLIA & REPASSES", '=ABS(SUMIFS(\'Transações Consolidadas\'!G:G, \'Transações Consolidadas\'!D:D, "Família & Pessoal", \'Transações Consolidadas\'!F:F, "Despesa"))', "F5", "G5"),
        ("RESULTADO DE CAIXA", "=C5-E5-G5", "H5", "I5"),
    ]

    for title, formula, c_lbl, c_val in cards:
        ws_dash[c_lbl] = title
        ws_dash[c_lbl].font = Font(name="Calibri", size=9, bold=True, color="595959")
        ws_dash[c_lbl].fill = fill_card_bg
        ws_dash[c_lbl].border = border_cell
        
        ws_dash[c_val] = formula
        ws_dash[c_val].font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
        ws_dash[c_val].fill = fill_card_bg
        ws_dash[c_val].border = border_cell
        ws_dash[c_val].number_format = '"R$" #,##0.00'

    ws_dash.column_dimensions['A'].width = 3
    ws_dash.column_dimensions['B'].width = 22
    ws_dash.column_dimensions['C'].width = 16
    ws_dash.column_dimensions['D'].width = 24
    ws_dash.column_dimensions['E'].width = 16
    ws_dash.column_dimensions['F'].width = 22
    ws_dash.column_dimensions['G'].width = 16
    ws_dash.column_dimensions['H'].width = 20
    ws_dash.column_dimensions['I'].width = 16

    wb.save("saida/FinanceAgent_Controle_Mensal.xlsx")
    print("Planilha 'saida/FinanceAgent_Controle_Mensal.xlsx' gerada com sucesso!")

if __name__ == "__main__":
    gerar_planilha()
